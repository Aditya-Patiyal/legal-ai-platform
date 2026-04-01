from __future__ import annotations

import re
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any

import docx
import fitz


@dataclass
class StructuredChunk:
    """A chunk with rich metadata for better retrieval."""
    text: str
    page_number: int | None = None
    section_heading: str | None = None
    clause_number: str | None = None
    chunk_index: int = 0
    char_start: int = 0
    char_end: int = 0
    
    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


SECTION_PATTERNS = [
    r"^(ARTICLE|Article|SECTION|Section|CLAUSE|Clause)\s*[IVXLCDM\d]+[.:]",
    r"^\d+\.\d*\s+[A-Z]",
    r"^[IVXLCDM]+\.\s+[A-Z]",
    r"^(SCHEDULE|Schedule|ANNEXURE|Annexure|APPENDIX|Appendix)\s*[A-Z\d]*",
    r"^(WHEREAS|NOW THEREFORE|IN WITNESS WHEREOF)",
    r"^(DEFINITIONS|TERM|TERMINATION|PAYMENT|CONFIDENTIALITY|INDEMNITY|LIABILITY|GOVERNING LAW|DISPUTE|ARBITRATION|FORCE MAJEURE|NOTICES|AMENDMENT|WAIVER|SEVERABILITY|ENTIRE AGREEMENT)",
]

CLAUSE_NUMBER_PATTERN = re.compile(r"^(\d+(?:\.\d+)*|[IVXLCDM]+|[a-z]\)|\([a-z]\)|\(\d+\))")


def detect_section_heading(text: str) -> str | None:
    """Detect if text starts with a section/clause heading."""
    first_line = text.split("\n")[0].strip()[:100]
    for pattern in SECTION_PATTERNS:
        if re.match(pattern, first_line, re.IGNORECASE):
            return first_line[:80]
    return None


def detect_clause_number(text: str) -> str | None:
    """Extract clause/section number from text."""
    first_line = text.split("\n")[0].strip()
    match = CLAUSE_NUMBER_PATTERN.match(first_line)
    if match:
        return match.group(1)
    return None


def extract_text_from_pdf(file_path: str) -> str:
    document = fitz.open(file_path)
    pages: list[str] = []
    for page in document:
        pages.append(page.get_text("text"))
    document.close()
    return "\n".join(pages).strip()


def extract_text_from_pdf_with_pages(file_path: str) -> list[tuple[int, str]]:
    """Extract text from PDF with page numbers."""
    document = fitz.open(file_path)
    pages: list[tuple[int, str]] = []
    for page_num, page in enumerate(document, start=1):
        text = page.get_text("text").strip()
        if text:
            pages.append((page_num, text))
    document.close()
    return pages


def extract_text_from_docx(file_path: str) -> str:
    document = docx.Document(file_path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n".join(paragraphs).strip()


def extract_text_from_docx_with_structure(file_path: str) -> list[tuple[int, str, str | None]]:
    """Extract text from DOCX with paragraph index and style info."""
    document = docx.Document(file_path)
    result: list[tuple[int, str, str | None]] = []
    for idx, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text:
            style_name = paragraph.style.name if paragraph.style else None
            result.append((idx + 1, text, style_name))
    return result


def extract_text(file_path: str) -> str:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(file_path)
    if suffix == ".docx":
        return extract_text_from_docx(file_path)
    raise ValueError("Unsupported file type")


def normalize_text(text: str) -> str:
    lines = [line.strip() for line in text.splitlines()]
    cleaned = [line for line in lines if line]
    return "\n".join(cleaned)


def chunk_text(text: str, chunk_size: int = 900, overlap: int = 150) -> list[str]:
    cleaned = normalize_text(text)
    if not cleaned:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(cleaned):
        end = start + chunk_size
        chunk = cleaned[start:end]
        chunks.append(chunk)
        if end >= len(cleaned):
            break
        start = max(end - overlap, 0)
    return chunks


def chunk_text_structured(
    file_path: str,
    chunk_size: int = 900,
    overlap: int = 150,
) -> list[StructuredChunk]:
    """
    Create structure-aware chunks with metadata for better retrieval.
    Preserves page numbers, section headings, and clause numbers.
    """
    suffix = Path(file_path).suffix.lower()
    chunks: list[StructuredChunk] = []
    
    if suffix == ".pdf":
        pages = extract_text_from_pdf_with_pages(file_path)
        chunk_index = 0
        
        for page_num, page_text in pages:
            cleaned = normalize_text(page_text)
            if not cleaned:
                continue
            
            start = 0
            while start < len(cleaned):
                end = min(start + chunk_size, len(cleaned))
                
                if end < len(cleaned):
                    for sep in ["\n\n", "\n", ". ", " "]:
                        last_sep = cleaned.rfind(sep, start, end)
                        if last_sep > start + chunk_size // 2:
                            end = last_sep + len(sep)
                            break
                
                chunk_text_content = cleaned[start:end].strip()
                if chunk_text_content:
                    chunks.append(StructuredChunk(
                        text=chunk_text_content,
                        page_number=page_num,
                        section_heading=detect_section_heading(chunk_text_content),
                        clause_number=detect_clause_number(chunk_text_content),
                        chunk_index=chunk_index,
                        char_start=start,
                        char_end=end,
                    ))
                    chunk_index += 1
                
                if end >= len(cleaned):
                    break
                start = max(end - overlap, start + 1)
    
    elif suffix == ".docx":
        paragraphs = extract_text_from_docx_with_structure(file_path)
        
        current_text = ""
        current_heading = None
        chunk_index = 0
        char_pos = 0
        
        for para_idx, para_text, style_name in paragraphs:
            if style_name and "heading" in style_name.lower():
                current_heading = para_text[:80]
            
            if len(current_text) + len(para_text) > chunk_size and current_text:
                chunks.append(StructuredChunk(
                    text=current_text.strip(),
                    page_number=None,
                    section_heading=current_heading or detect_section_heading(current_text),
                    clause_number=detect_clause_number(current_text),
                    chunk_index=chunk_index,
                    char_start=char_pos,
                    char_end=char_pos + len(current_text),
                ))
                chunk_index += 1
                char_pos += len(current_text)
                
                overlap_text = current_text[-overlap:] if len(current_text) > overlap else current_text
                current_text = overlap_text + "\n" + para_text
            else:
                current_text += ("\n" if current_text else "") + para_text
        
        if current_text.strip():
            chunks.append(StructuredChunk(
                text=current_text.strip(),
                page_number=None,
                section_heading=current_heading or detect_section_heading(current_text),
                clause_number=detect_clause_number(current_text),
                chunk_index=chunk_index,
                char_start=char_pos,
                char_end=char_pos + len(current_text),
            ))
    
    else:
        raise ValueError("Unsupported file type")
    
    return chunks
